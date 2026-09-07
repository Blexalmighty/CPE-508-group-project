from pathlib import Path

from docx import Document
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas

DOC_PATH = Path(__file__).resolve().parent
SOURCE = DOC_PATH / 'CODEBASE_DOCUMENTATION.md'
WORD_OUT = DOC_PATH / 'OncoPredict_Codebase_Documentation.docx'
PDF_OUT = DOC_PATH / 'OncoPredict_Codebase_Documentation.pdf'

text = SOURCE.read_text(encoding='utf-8')

# Create Word document
word_doc = Document()
for line in text.splitlines():
    if not line.strip():
        word_doc.add_paragraph('')
        continue
    word_doc.add_paragraph(line)
word_doc.save(WORD_OUT)

# Create PDF
pdf = canvas.Canvas(str(PDF_OUT), pagesize=letter)
pdf.setTitle('OncoPredict CDSS Codebase Documentation')
pdf.setAuthor('GitHub Copilot')

lines = text.splitlines()
page_width, page_height = letter
left_margin = 50
top_margin = 760
line_height = 14
max_chars_per_line = 100

for index, raw_line in enumerate(lines):
    y = top_margin - (index % 50) * line_height
    if index % 50 == 0 and index != 0:
        pdf.showPage()
        y = top_margin

    cleaned = raw_line[:max_chars_per_line]
    pdf.drawString(left_margin, y, cleaned)

pdf.save()

print(f'Created: {WORD_OUT}')
print(f'Created: {PDF_OUT}')
